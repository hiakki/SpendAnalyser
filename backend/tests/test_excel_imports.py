"""Real synthetic XLS/XLSX files exercise readers and multipart ingestion."""

import atexit
from datetime import date, datetime
from io import BytesIO
import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

_scratch = tempfile.TemporaryDirectory(prefix=".test-excel-", dir=Path(__file__).resolve().parents[1])
atexit.register(_scratch.cleanup)
os.environ["SPENDA_DB_URL"] = "sqlite://"
os.environ["SPENDA_UPLOAD_DIR"] = _scratch.name
os.environ["SPENDA_LLM_ENABLED"] = "false"

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
import xlwt

from app import models
from app.api import upload
from app.db import Base, get_db
from app.parsers.xls_parser import parse_xls


def workbook_bytes(kind: str, *, preamble: int = 0, notice: bool = False) -> bytes:
    rows = [[f"Synthetic account metadata {index}"] for index in range(preamble)]
    if notice:
        rows.append(["Transaction date description amount debit credit withdrawal deposit balance reference details"])
    rows.extend([
        ["Transaction Date", "Description", "Debit", "Credit", "Balance"],
        [datetime(2026, 9, 15), "Synthetic grocery", 125.5, None, 874.5],
        [datetime(2026, 9, 16), "Synthetic salary", None, 1000, 1874.5],
    ])
    output = BytesIO()
    if kind == "xlsx":
        workbook = Workbook()
        workbook.active.title = "Cover"
        workbook.active.append(["Synthetic statement information"])
        sheet = workbook.create_sheet("Transactions")
        for row in rows:
            sheet.append(row)
        workbook.save(output)
        workbook.close()
    else:
        workbook = xlwt.Workbook()
        workbook.add_sheet("Cover").write(0, 0, "Synthetic statement information")
        sheet = workbook.add_sheet("Transactions")
        date_style = xlwt.easyxf(num_format_str="DD/MM/YYYY")
        for row_number, row in enumerate(rows):
            for column_number, value in enumerate(row):
                if value is not None:
                    if isinstance(value, datetime):
                        sheet.write(row_number, column_number, value, date_style)
                    else:
                        sheet.write(row_number, column_number, value)
        workbook.save(output)
    return output.getvalue()


class ExcelParserTests(unittest.TestCase):
    def parse_fixture(self, kind, *, suffix=None, **kwargs):
        path = Path(_scratch.name) / f"synthetic-{kind}.{suffix or kind}"
        path.write_bytes(workbook_bytes(kind, **kwargs))
        return parse_xls(path)

    def assert_transactions(self, result, kind):
        self.assertEqual(result.detected_format, kind)
        self.assertEqual([row.amount for row in result.rows], [-125.5, 1000])
        self.assertEqual([row.direction for row in result.rows], ["debit", "credit"])
        self.assertEqual(result.period_start, date(2026, 9, 15))
        self.assertEqual(result.period_end, date(2026, 9, 16))

    def test_real_xlsx_native_dates_and_blank_amount_cells(self):
        self.assert_transactions(self.parse_fixture("xlsx"), "xlsx")

    def test_real_binary_xls_native_dates_and_blank_amount_cells(self):
        self.assert_transactions(self.parse_fixture("xls"), "xls")

    def test_metadata_notice_cannot_outscore_real_headers(self):
        for kind in ("xls", "xlsx"):
            with self.subTest(kind=kind):
                self.assert_transactions(self.parse_fixture(kind, notice=True), kind)

    def test_long_metadata_preamble_and_cover_sheet(self):
        for kind in ("xls", "xlsx"):
            with self.subTest(kind=kind):
                result = self.parse_fixture(kind, preamble=30, notice=True)
                self.assert_transactions(result, kind)
                self.assertTrue(any("Cover" in warning for warning in result.warnings))

    def test_workbook_content_selects_engine_when_extension_is_wrong(self):
        for kind, suffix in (("xlsx", "xls"), ("xls", "xlsx")):
            with self.subTest(kind=kind, suffix=suffix):
                self.assert_transactions(self.parse_fixture(kind, suffix=suffix), kind)


class ExcelUploadTests(unittest.TestCase):
    def setUp(self):
        settings_patch = patch.object(upload, "get_settings", return_value=SimpleNamespace(upload_dir=_scratch.name))
        settings_patch.start()
        self.addCleanup(settings_patch.stop)
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine, expire_on_commit=False, autoflush=False)
        self.db.add(models.Category(name="Other"))
        self.db.commit()
        app = FastAPI()
        app.include_router(upload.router)
        app.dependency_overrides[get_db] = lambda: self.db
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.db.close()
        self.engine.dispose()

    def test_both_formats_upload_and_reupload_without_duplicates(self):
        for kind in ("xls", "xlsx"):
            with self.subTest(kind=kind):
                account = models.Account(name=f"Synthetic {kind} account", kind="bank")
                self.db.add(account)
                self.db.commit()
                payload = workbook_bytes(kind, preamble=30, notice=True)
                response = self.client.post("/upload", data={"account_id": account.id}, files={"file": (f"synthetic.{kind}", payload)})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["inserted"], 2)
                self.assertEqual(response.json()["statement"]["detected_format"], kind)
                repeated = self.client.post("/upload", data={"account_id": account.id}, files={"file": (f"synthetic.{kind}", payload)})
                self.assertEqual(repeated.status_code, 200, repeated.text)
                self.assertEqual((repeated.json()["inserted"], repeated.json()["duplicates"]), (0, 2))
                amounts = [row.raw_amount for row in self.db.query(models.Transaction).filter_by(account_id=account.id).order_by(models.Transaction.posted_at)]
                self.assertEqual(amounts, [-125.5, 1000])


if __name__ == "__main__":
    unittest.main()
